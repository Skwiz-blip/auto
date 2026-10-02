"""Remplit directement le fichier SharePoint (MODELE SHAREPOINT VC.xlsx) avec les dossiers VALIDÉS.

Une ligne par dossier validé, ajoutée à la suite des lignes déjà présentes.
Le texte manuscrit mal lu est reporté tel quel (« ? » à la place des caractères illisibles) :
le dossier est validé, seule la saisie reste à compléter à la main.
"""
import json
import re
import shutil
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

import chemins
import referentiel
from bench_dossiers import SCHEMA, conform, nom_valide
from qr_officiel import numero_benin

MODELE = chemins.MODELE
A_COMPLETER = PatternFill("solid", fgColor="FFD8A8")  # orange clair


PRINCIPAL_MAX, SOUS_COMPTES_MAX = 10, 20  # au-delà : chiffre mal lu, valeur par défaut


def ligne(journal: dict, date_transmission: str) -> list:
    """Imprimé d'abord (RCCM, IFU), manuscrit ensuite : l'écriture de la fiche est la source la
    moins sûre. Ville et département sont ramenés au référentiel des communes du Bénin."""
    # conform : complète les champs absents des résultats plus anciens (ex. nombre_head)
    e = conform(journal["extraction"], SCHEMA)
    fiche, rccm, ifu, pid = e["fiche"], e["rccm"], e["ifu"], e["piece_identite"]
    structure = next((v for v in (rccm["enseigne"], rccm["nom_commercial"],
                                   ifu["nom_etablissement"]) if nom_valide(v)),
                     fiche["nom_structure"]).strip()
    # promoteur : nom et prénoms du RCCM (à défaut, ceux de la pièce d'identité)
    nom = (rccm["nom"] or pid["nom"]).strip().upper()
    prenoms = (rccm["prenoms"] or pid["prenoms"]).strip().upper()
    # contact : numéro imprimé sur le RCCM, puis la CIP, l'IFU, celui écrit sur la fiche ;
    # au format KIK « 229 » + numéro (2290197155835)
    numero = (numero_benin(rccm["telephone"]) or numero_benin(pid["telephone"])
              or numero_benin(ifu["telephone"]) or numero_benin(fiche["telephone"]))
    contact = ("229" + numero) if numero else re.sub(r"[^\d?]", "", fiche["telephone"])
    ville = referentiel.commune(fiche["ville"])
    return [
        date_transmission,                                                   # A
        structure,                                                           # B STRUCTURE
        referentiel.nombre_comptes(fiche["nombre_head"], 1, PRINCIPAL_MAX),             # C
        referentiel.nombre_comptes(fiche["nombre_sous_comptes"], 0, SOUS_COMPTES_MAX),  # D
        nom, prenoms, contact,                                               # E-G PROMOTEUR
        nom, prenoms, contact,                                               # H-J GESTIONNAIRE
        fiche["secteur"].strip(),                                            # K SECTEUR
        referentiel.departement(fiche["departement"], ville, fiche["commercial"]),       # L
        ville or fiche["ville"].strip(),                                     # M VILLE/COMMUNE
        fiche["quartier"].strip(),                                           # N SITUATION GEO.
        referentiel.commercial(fiche["commercial"]),  # O : nom de la liste officielle, ou vide
        "",                                                                  # P (laissé vide)
    ]


def _derniere_ligne(feuille) -> int:
    """Dernière ligne réellement remplie (openpyxl compte parfois des lignes vides mises en forme)."""
    for n in range(feuille.max_row, 1, -1):
        if any(feuille.cell(row=n, column=c).value not in (None, "") for c in range(1, 17)):
            return n
    return 1


def remplir(sortie: Path, journaux: list[dict], date_transmission: str = "") -> Path | None:
    """Ajoute les dossiers VALIDÉS à la suite des lignes de MODELE SHAREPOINT VC.xlsx.

    Un même lancement n'est jamais ajouté deux fois : les dossiers déjà reportés sont notés
    dans <sortie>/sharepoint_ajoutes.txt. Une copie du modèle avant écriture est gardée dans
    <sortie>/MODELE SHAREPOINT VC (avant).xlsx. None si rien de nouveau à ajouter."""
    registre = sortie / "sharepoint_ajoutes.txt"
    deja = set(registre.read_text(encoding="utf-8").splitlines()) if registre.exists() else set()
    valides = sorted((j for j in journaux
                      if (j.get("decision") or {}).get("statut") == "VALIDÉ"
                      and j["dossier"] not in deja),
                     key=lambda j: j["dossier"])
    if not valides or not MODELE.exists():
        return None
    date_transmission = date_transmission or date.today().strftime("%d/%m/%Y")
    classeur = load_workbook(MODELE)
    feuille = classeur.worksheets[0]
    depart = _derniere_ligne(feuille) + 1
    for n, journal in enumerate(valides, start=depart):
        for col, valeur in enumerate(ligne(journal, date_transmission), start=1):
            cellule = feuille.cell(row=n, column=col, value=valeur)
            # numéros en texte (garde le 0 de « 01… ») ; principal et sous-comptes en nombres
            cellule.number_format = "General" if isinstance(valeur, int) else "@"
            # à compléter à la main : caractère illisible, contact non conforme
            if ("?" in str(valeur) or (col in (7, 10) and not re.fullmatch(r"22901\d{8}", valeur))
                    or (col in (12, 15) and not valeur)):
                cellule.fill = A_COMPLETER
    shutil.copy2(MODELE, sortie / f"{MODELE.stem} (avant).xlsx")
    classeur.save(MODELE)  # PermissionError si le fichier est ouvert dans Excel
    with open(registre, "a", encoding="utf-8") as f:
        f.write("".join(j["dossier"] + "\n" for j in valides))
    return MODELE


def remplir_depuis(sortie: Path) -> Path | None:
    journaux = [json.loads(l) for l in (sortie / "resultats.jsonl").read_text(
        encoding="utf-8").splitlines() if l.strip()]
    return remplir(sortie, journaux)


if __name__ == "__main__":
    # python sharepoint.py <dossier de sortie d'un lancement>
    print(remplir_depuis(Path(sys.argv[1])))
