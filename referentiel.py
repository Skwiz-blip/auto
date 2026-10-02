"""Référentiel géographique du Bénin et normalisation des champs de la fiche manuscrite.

La fiche est remplie à la main : « Natitingui », « AB-Calavi », « ZALI »… On ramène la ville à
l'une des 77 communes (ou à la commune d'un arrondissement connu) et on en déduit le
département, ce qui évite aussi les incohérences du type « ZOU » + « COTONOU ».
"""
import re
import unicodedata
from difflib import SequenceMatcher

import chemins

COMMUNES = {
    "ALIBORI": ["Banikoara", "Gogounou", "Kandi", "Karimama", "Malanville", "Ségbana"],
    "ATACORA": ["Boukoumbé", "Cobly", "Kérou", "Kouandé", "Matéri", "Natitingou", "Péhunco",
                "Tanguiéta", "Toucountouna"],
    "ATLANTIQUE": ["Abomey-Calavi", "Allada", "Kpomassè", "Ouidah", "Sô-Ava", "Toffo",
                   "Tori-Bossito", "Zè"],
    "BORGOU": ["Bembèrèkè", "Kalalé", "N'Dali", "Nikki", "Parakou", "Pèrèrè", "Sinendé",
               "Tchaourou"],
    "COLLINES": ["Bantè", "Dassa-Zoumè", "Glazoué", "Ouèssè", "Savalou", "Savè"],
    "COUFFO": ["Aplahoué", "Djakotomey", "Dogbo", "Klouékanmè", "Lalo", "Toviklin"],
    "DONGA": ["Bassila", "Copargo", "Djougou", "Ouaké"],
    "LITTORAL": ["Cotonou"],
    "MONO": ["Athiémé", "Bopa", "Comè", "Grand-Popo", "Houéyogbé", "Lokossa"],
    "OUEME": ["Adjarra", "Adjohoun", "Aguégués", "Akpro-Missérété", "Avrankou", "Bonou",
              "Dangbo", "Porto-Novo", "Sèmè-Kpodji"],
    "PLATEAU": ["Adja-Ouèrè", "Ifangni", "Kétou", "Pobè", "Sakété"],
    "ZOU": ["Abomey", "Agbangnizoun", "Bohicon", "Covè", "Djidja", "Ouinhi", "Zagnanado",
            "Za-Kpota", "Zogbodomey"],
}
DEPARTEMENTS = list(COMMUNES) 
LOCALITES = {
    "CALAVI": "Abomey-Calavi", "ABCALAVI": "Abomey-Calavi", "AKASSATO": "Abomey-Calavi",
    "GODOMEY": "Abomey-Calavi", "OUEDO": "Abomey-Calavi", "HEVIE": "Abomey-Calavi",
    "TOGBA": "Abomey-Calavi", "ZINVIE": "Abomey-Calavi", "GLODJIGBE": "Abomey-Calavi",
    "KPANROUN": "Abomey-Calavi", "COCOTOMEY": "Abomey-Calavi", "COCODJI": "Abomey-Calavi",
    "COCOCODJI": "Abomey-Calavi", "ZOGBADJE": "Abomey-Calavi", "PAHOU": "Ouidah",
    "SAVI": "Ouidah", "HOUEGBO": "Toffo", "SEKOU": "Allada", "ATTOGON": "Allada",
    "TORI": "Tori-Bossito", "SEME": "Sèmè-Kpodji", "DJEFFA": "Sèmè-Kpodji",
    "EKPE": "Sèmè-Kpodji", "PORTONOVO": "Porto-Novo", "DASSA": "Dassa-Zoumè",
    "DOGBO": "Dogbo", "NATITINGOU": "Natitingou",
    "AKOSSAVIE": "Abomey-Calavi", "COCOCODJI": "Abomey-Calavi", "GOLODJIGBE": "Abomey-Calavi",
    "DJADJO": "Abomey-Calavi", "BIDOSSESSI": "Abomey-Calavi", "KPOSSIDJA": "Abomey-Calavi",
    "KILIBO": "Ouèssè", "AGBANGNIZOUN": "Agbangnizoun", "BOHIQUE": "Bohicon",
    "GOUNLI": "Covè", "AZOVE": "Aplahoué",
}

def _charger_commerciaux() -> dict[str, str]: 
    try:
        from openpyxl import load_workbook
        feuille = load_workbook(chemins.COMMERCIAUX, read_only=True).worksheets[0]
        lignes = list(feuille.iter_rows(values_only=True))
    except Exception:  # fichier absent, ouvert ailleurs ou illisible
        return {}
    entete = [_cle(str(v or "")) for v in lignes[0]] if lignes else []
    col_nom = entete.index("BDP") if "BDP" in entete else len(entete) - 1
    col_dept = next((i for i, e in enumerate(entete) if e.startswith("DEPARTEMENT")), None)
    liste = {}
    for ligne in lignes[1:]:
        if col_nom < len(ligne) and ligne[col_nom] and str(ligne[col_nom]).strip():
            dept = str(ligne[col_dept] or "").strip() if col_dept is not None else ""
            liste.setdefault(" ".join(str(ligne[col_nom]).split()), dept)
    return liste


_COMMUNE_DEPT = {c: d for d, liste in COMMUNES.items() for c in liste}
_NOMS_DEPT = {"ALIBORI": "Alibori", "ATACORA": "Atacora", "ATLANTIQUE": "Atlantique",
              "BORGOU": "Borgou", "COLLINES": "Collines", "COUFFO": "Couffo",
              "DONGA": "Donga", "LITTORAL": "Littoral", "MONO": "Mono", "OUEME": "Ouémé",
              "PLATEAU": "Plateau", "ZOU": "Zou"}


def _cle(texte: str) -> str:
    t = unicodedata.normalize("NFKD", texte or "")
    return re.sub(r"[^A-Z]", "", "".join(c for c in t if not unicodedata.combining(c)).upper())


def _plus_proche(cle: str, choix: dict[str, str], seuil: float) -> str:
    if not cle:
        return ""
    meilleur = max(choix, key=lambda c: SequenceMatcher(None, cle, c).ratio())
    return choix[meilleur] if SequenceMatcher(None, cle, meilleur).ratio() >= seuil else ""


# liste officielle : seuls ces noms sont reconnus (fiche, fichier SharePoint, application)
DEPT_COMMERCIAL = _charger_commerciaux()
COMMERCIAUX = list(DEPT_COMMERCIAL)
# départements : uniquement ceux du fichier officiel, rien de plus
DEPARTEMENTS_KIK = list(dict.fromkeys(d for d in DEPT_COMMERCIAL.values() if d))


def commune(ville_lue: str) -> str: 
    cle = _cle(ville_lue)
    if not cle:
        return ""
    choix = {_cle(c): c for c in _COMMUNE_DEPT} | LOCALITES
    if cle in choix:
        return choix[cle]
    # nom entier proche d'abord (« ABOMEY CALAC? » -> Abomey-Calavi, pas Abomey)
    if proche := _plus_proche(cle, choix, 0.8):
        return proche
    # « ABOMEY CALAVI AKASSATO » : une localité connue contenue dans le texte
    for loc in sorted(choix, key=len, reverse=True):
        if len(loc) >= 5 and loc in cle:
            return choix[loc]
    return _plus_proche(cle, choix, 0.7)


def departement(dept_lu: str, commune_trouvee: str = "", commercial_lu: str = "") -> str:
    choix = {_cle(d): d for d in DEPARTEMENTS_KIK}
    # commune reconnue : son département, s'il fait partie de ceux du fichier (Cotonou ->
    # Littoral n'y est pas : on passe au département du commercial, puis à l'écriture)
    if commune_trouvee in _COMMUNE_DEPT and _COMMUNE_DEPT[commune_trouvee] in choix:
        return choix[_COMMUNE_DEPT[commune_trouvee]]
    # chaque commercial (BDP) est rattaché à un département dans le fichier
    if commercial_lu and (dept := DEPT_COMMERCIAL.get(commercial(commercial_lu), "")):
        return dept
    cle = _cle(dept_lu)
    if len(cle) >= 3 and choix:
        # abréviation : « ATL », « COLL »
        debut = [d for c, d in choix.items() if c.startswith(cle)]
        if len(debut) == 1:
            return debut[0]
        if trouve := _plus_proche(cle, choix, 0.6):
            return trouve
    return ""


def _ressemblance(lu: str, nom: str) -> float: 
    mots_nom = [_cle(m) for m in nom.split() if len(_cle(m)) >= 2]
    mots_lus = [_cle(m) for m in re.split(r"[\s.\-]+", lu) if len(_cle(m)) >= 2]
    if not mots_nom or not mots_lus:
        return 0.0
    entier = max(SequenceMatcher(None, _cle(lu), "".join(ordre)).ratio()
                 for ordre in (mots_nom, mots_nom[::-1]))

    mot_a_mot = sum(len(m) * max(SequenceMatcher(None, m, l).ratio() for l in mots_lus)
                    for m in mots_nom) / sum(len(m) for m in mots_nom)
    return max(entier, mot_a_mot)


def commercial(lu: str) -> str: 
    texte = re.sub(r"(?i)^\s*demand[ée]e? par\s*:?\s*", "", lu or "").replace("?", "").strip()
    if len(_cle(texte)) < 3 or not COMMERCIAUX:
        return ""
    scores = sorted(((_ressemblance(texte, nom), nom) for nom in COMMERCIAUX), reverse=True)
    meilleur, nom = scores[0]
    second = scores[1][0] if len(scores) > 1 else 0.0
    return nom if meilleur >= 0.85 or (meilleur >= 0.6 and meilleur - second >= 0.08) else ""


def nombre_comptes(valeur: str, defaut: int) -> int:
    # 0 ou 1 seulement : un 1 écrit -> 1 ; que des zéros -> 0 ; vide ou autre chiffre (zéro
    # stylisé lu « 5 », « 2 »…) -> défaut (1 pour le head, 0 pour les sous-comptes)
    chiffres = re.sub(r"\D", "", valeur or "")
    if "1" in chiffres:
        return 1
    if chiffres and set(chiffres) == {"0"}:
        return 0
    return defaut
